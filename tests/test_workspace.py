from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from PIL import Image
from reportlab.pdfgen import canvas

from extract_materials import build_bundle, discover_files
from render_outputs import render_course


def test_discovery_ignores_generated_directories(course_dir: Path) -> None:
    source = course_dir / "讲义.docx"
    document = Document()
    document.add_paragraph("课程内容")
    document.save(source)
    ignored = course_dir / ".tutor" / "cache"
    ignored.mkdir(parents=True)
    document.save(ignored / "ignored.docx")
    assert discover_files(course_dir) == [source]


def test_bundle_and_manifest(course_dir: Path) -> None:
    source = course_dir / "历年真题.docx"
    document = Document()
    document.add_paragraph("第一题：解释时间复杂度。")
    document.save(source)

    bundle = build_bundle(course_dir)
    assert bundle["manifest"]["successful_files"] == 1
    assert bundle["manifest"]["files"][0]["signal_strength"] == 4
    assert "时间复杂度" in bundle["merged_text"]
    assert (course_dir / ".tutor" / "manifest.json").exists()
    assert (course_dir / ".tutor" / "visual_catalog.json").exists()


def test_render_markdown_outputs(course_dir: Path) -> None:
    output_dir = course_dir / "outputs"
    output_dir.mkdir()
    source_image = course_dir / "diagram.png"
    Image.new("RGB", (80, 60), "white").save(source_image)
    (output_dir / "review-guide.md").write_text(
        "# 第一章\n\n**核心概念**：栈是后进先出。\n\n"
        "- push\n- pop\n\n"
        "| 结构 | 规则 |\n|---|---|\n| 栈 | 后进先出 |\n\n"
        "行内公式 \\(x^2\\)。\n\n"
        "\\[\n\\lim_{x \\to 1}\\frac{x^2-1}{x-1}=2\n\\]\n\n"
        "![资料原图：函数示意](diagram.png)",
        encoding="utf-8",
    )
    destination = render_course(course_dir)
    html = destination.read_text(encoding="utf-8")
    assert destination.name == "review-site.html"
    assert "栈是后进先出" in html
    assert "<strong>核心概念</strong>" in html
    assert "<table>" in html
    assert "<td>后进先出</td>" in html
    assert "assets/katex/katex.min.js" in html
    assert '<div class="math-block">' in html
    assert html.count(r"\lim_{x \to 1}\frac{x^2-1}{x-1}=2") == 1
    assert (output_dir / "assets" / "katex" / "katex.min.css").exists()
    assert "assets/visuals/" in html
    assert list((output_dir / "assets" / "visuals").glob("*-diagram.png"))


def test_render_falls_back_to_extraction_bundle(course_dir: Path) -> None:
    tutor_dir = course_dir / ".tutor"
    tutor_dir.mkdir()
    (tutor_dir / "extraction_bundle.json").write_text(
        json.dumps({"merged_text": "# 提取内容\n\n测试"}, ensure_ascii=False),
        encoding="utf-8",
    )
    destination = render_course(course_dir)
    assert destination.exists()
    assert (course_dir / "outputs" / "review-guide.md").exists()


def test_visual_catalog_tracks_standalone_image(course_dir: Path) -> None:
    image = course_dir / "函数图.png"
    Image.new("RGB", (120, 80), "white").save(image)
    bundle = build_bundle(course_dir)
    catalog = json.loads(
        (course_dir / ".tutor" / "visual_catalog.json").read_text(encoding="utf-8")
    )
    assert bundle["manifest"]["visual_count"] == 1
    visual = catalog["visuals"][0]
    assert visual["source_path"] == "函数图.png"
    assert visual["locator"] == "standalone"
    assert visual["provenance"] == "资料原图"
    assert visual["multimodal_review_status"] == "pending"


def test_visual_catalog_creates_pdf_page_snapshot(course_dir: Path) -> None:
    source = course_dir / "函数讲义.pdf"
    pdf = canvas.Canvas(str(source))
    pdf.drawString(72, 760, "Function graph and limit")
    pdf.line(72, 700, 300, 500)
    pdf.save()
    bundle = build_bundle(course_dir)
    catalog = json.loads(
        (course_dir / ".tutor" / "visual_catalog.json").read_text(encoding="utf-8")
    )
    snapshots = [
        visual
        for visual in catalog["visuals"]
        if visual["kind"] == "page-snapshot"
    ]
    assert bundle["manifest"]["visual_count"] >= 1
    assert snapshots[0]["locator"] == "page:1:snapshot"
    assert (course_dir / snapshots[0]["path"]).exists()
