from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document
from PIL import Image
from pptx import Presentation
from reportlab.pdfgen import canvas

import extract_pdf as pdf_module
from extract_docx import build_text_summary as docx_summary
from extract_docx import extract_docx
from extract_pdf import build_text_summary as pdf_summary
from extract_pdf import extract_pdf
from extract_pptx import build_text_summary as pptx_summary
from extract_pptx import extract_pptx


def test_docx_extraction(tmp_path: Path) -> None:
    path = tmp_path / "notes.docx"
    document = Document()
    document.add_heading("第一章", level=1)
    document.add_paragraph("时间复杂度描述算法增长趋势。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "概念"
    table.cell(0, 1).text = "含义"
    table.cell(1, 0).text = "O(n)"
    table.cell(1, 1).text = "线性"
    document.save(path)

    result = extract_docx(str(path))
    assert result["sections"][0]["heading"] == "第一章"
    assert "时间复杂度" in docx_summary(result)
    assert result["tables"][0][1][0] == "O(n)"


def test_pptx_extraction(tmp_path: Path) -> None:
    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "栈和队列"
    slide.placeholders[1].text = "栈是后进先出"
    presentation.save(path)

    result = extract_pptx(str(path))
    assert result["total_slides"] == 1
    assert "栈和队列" in pptx_summary(result)
    assert "后进先出" in result["slides"][0]["text"]


def test_text_pdf_does_not_need_ocr(tmp_path: Path) -> None:
    path = tmp_path / "text.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 760, "Convolutional neural networks use filters for feature extraction.")
    pdf.save()

    result = extract_pdf(str(path), minimum_native_characters=10)
    assert result["total_pages"] == 1
    assert result["pages"][0]["ocr_attempted"] is False
    assert "Convolutional" in pdf_summary(result)


def test_scanned_pdf_uses_ocr_fallback(tmp_path: Path, monkeypatch) -> None:
    image_path = tmp_path / "scan.png"
    Image.new("RGB", (300, 120), "white").save(image_path)
    pdf_path = tmp_path / "scan.pdf"
    document = fitz.open()
    page = document.new_page(width=300, height=120)
    page.insert_image(page.rect, filename=str(image_path))
    document.save(pdf_path)
    document.close()

    monkeypatch.setattr(
        pdf_module,
        "ocr_image",
        lambda image, languages: {
            "text": "扫描课程内容",
            "confidence": 91.5,
        },
    )
    result = extract_pdf(str(pdf_path))
    page_result = result["pages"][0]
    assert page_result["ocr_attempted"] is True
    assert page_result["ocr_used"] is True
    assert page_result["text"] == "扫描课程内容"


def test_corrupted_files_raise(tmp_path: Path) -> None:
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"not a pdf")
    try:
        extract_pdf(str(path))
    except Exception:
        pass
    else:
        raise AssertionError("corrupted PDF should fail")
