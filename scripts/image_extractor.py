"""Extract embedded images and OCR standalone image files."""

from __future__ import annotations

import io
import os
import zipfile
from pathlib import Path

import fitz
from PIL import Image
from pptx import Presentation

from ocr import DEFAULT_LANGUAGES, ocr_image


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}


def _ensure_output(output_dir: str | Path) -> Path:
    path = Path(output_dir)
    path.mkdir(parents=True, exist_ok=True)
    return path


def extract_from_pptx(filepath: str, output_dir: str) -> list[str]:
    output = _ensure_output(output_dir)
    paths = []
    presentation = Presentation(filepath)
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes, start=1):
            if getattr(shape, "shape_type", None) != 13:
                continue
            image = shape.image
            extension = image.ext or "bin"
            path = output / f"slide-{slide_index}-image-{shape_index}.{extension}"
            path.write_bytes(image.blob)
            paths.append(str(path))
    return paths


def extract_from_docx(filepath: str, output_dir: str) -> list[str]:
    output = _ensure_output(output_dir)
    paths = []
    with zipfile.ZipFile(filepath) as archive:
        for member in archive.namelist():
            if not member.startswith("word/media/") or member.endswith("/"):
                continue
            data = archive.read(member)
            name = Path(member).name
            try:
                image = Image.open(io.BytesIO(data))
                path = output / f"{Path(name).stem}.png"
                image.save(path, "PNG")
            except Exception:
                path = output / name
                path.write_bytes(data)
            paths.append(str(path))
    return paths


def extract_from_pdf(filepath: str, output_dir: str) -> list[str]:
    output = _ensure_output(output_dir)
    paths = []
    document = fitz.open(filepath)
    try:
        for page_index, page in enumerate(document, start=1):
            for image_index, image_info in enumerate(page.get_images(full=True), start=1):
                extracted = document.extract_image(image_info[0])
                path = output / (
                    f"page-{page_index}-image-{image_index}.{extracted['ext']}"
                )
                path.write_bytes(extracted["image"])
                paths.append(str(path))
    finally:
        document.close()
    return paths


def extract_image_text(
    filepath: str, languages: str = DEFAULT_LANGUAGES
) -> dict[str, object]:
    result = ocr_image(filepath, languages=languages)
    return {
        "filename": os.path.basename(filepath),
        "text": result["text"],
        "ocr_confidence": result["confidence"],
        "ocr_used": True,
    }
