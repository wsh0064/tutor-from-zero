"""Extract slide text, notes, tables, and image counts from PPTX files."""

from __future__ import annotations

import os
from typing import Any

from pptx import Presentation


def _shape_text(shape: Any) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(
        paragraph.text.strip()
        for paragraph in shape.text_frame.paragraphs
        if paragraph.text.strip()
    )


def extract_pptx(filepath: str) -> dict[str, Any]:
    presentation = Presentation(filepath)
    slides = []
    for number, slide in enumerate(presentation.slides, start=1):
        title = ""
        texts: list[str] = []
        tables: list[list[list[str]]] = []
        image_count = 0
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                tables.append(
                    [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                )
            elif getattr(shape, "shape_type", None) == 13:
                image_count += 1
            else:
                text = _shape_text(shape)
                if not text:
                    continue
                if getattr(shape, "is_placeholder", False):
                    try:
                        if shape.placeholder_format.type == 1:
                            title = text
                            continue
                    except Exception:
                        pass
                texts.append(text)
        notes = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            pass
        slides.append(
            {
                "number": number,
                "title": title,
                "text": "\n".join(texts),
                "tables": tables,
                "image_count": image_count,
                "notes": notes,
            }
        )
    return {
        "filename": os.path.basename(filepath),
        "total_slides": len(slides),
        "slides": slides,
    }


def build_text_summary(data: dict[str, Any]) -> str:
    lines = [f"# {data['filename']}", f"共 {data['total_slides']} 张幻灯片"]
    for slide in data["slides"]:
        lines.append(f"\n## 幻灯片 {slide['number']}")
        if slide["title"]:
            lines.append(f"### {slide['title']}")
        if slide["text"]:
            lines.append(slide["text"])
        for table in slide["tables"]:
            if not table:
                continue
            width = max(len(row) for row in table)
            rows = [row + [""] * (width - len(row)) for row in table]
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("| " + " | ".join(["---"] * width) + " |")
            lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
        if slide["notes"]:
            lines.append(f"> 讲者备注：{slide['notes']}")
        if slide["image_count"]:
            lines.append(f"> 本页包含 {slide['image_count']} 张图片，可按需进行 OCR。")
    return "\n".join(lines)
